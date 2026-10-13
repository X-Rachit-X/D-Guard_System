"""
scripts/build_slides.py - The project presentation (18 slides, 16:9) with speaker notes.

    python scripts/build_slides.py      ->  reports/DERMA-GUARD_Presentation.pptx

Uses the figures from paper/figures/ (scripts/make_figures*.py) and, if present, the app
screenshots in reports/screens/. Every number on the slides is a final measured result.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "paper" / "figures"
SCR = ROOT / "reports" / "screens"
OUT = ROOT / "reports" / "DERMA-GUARD_Presentation.pptx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xC0, 0x50, 0x4D)
GREEN = RGBColor(0x1B, 0x7F, 0x3B)
AMBER = RGBColor(0xB5, 0x47, 0x08)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xEE, 0xF2, 0xFF)
GOLD = RGBColor(0xF2, 0xB7, 0x05)
C_HIGH = RGBColor(0x1B, 0x7F, 0x3B)
C_LOW = RGBColor(0x2E, 0x75, 0xB6)
C_ASK = RGBColor(0xC5, 0x7B, 0x00)
C_DEFER = RGBColor(0xB0, 0x3A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
N = {"i": 0}


def text(slide, x, y, w, h, lines, size=18, color=GREY, bold=False, align=PP_ALIGN.LEFT,
         bullet=False, spacing=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        parts = ln if isinstance(ln, list) else [(ln, bold, color)]
        if bullet:
            r = p.add_run()
            r.text = "•  "
            r.font.size, r.font.color.rgb, r.font.name = Pt(size), NAVY, FONT
        for t, b, c in parts:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), b, c, FONT
    return tb


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.adjustments[0] = 0.08
    return s


def card(slide, x, y, w, h, big, label, sub="", color=NAVY, big_size=34):
    """KPI card: coloured top strip, a big number, a label and a short note."""
    rect(slide, x, y, w, h, WHITE, line=RGBColor(0xD0, 0xD7, 0xE5))
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.12))
    top.fill.solid()
    top.fill.fore_color.rgb = color
    top.line.fill.background()
    big_h = Pt(big_size * 1.5)
    text(slide, x + Inches(0.1), y + Inches(0.18), w - Inches(0.2), big_h, big, big_size, color,
         True, PP_ALIGN.CENTER, spacing=0)
    text(slide, x + Inches(0.1), y + Inches(0.2) + big_h, w - Inches(0.2), Inches(0.45),
         label, 15, NAVY, True, PP_ALIGN.CENTER, spacing=0)
    if sub:
        text(slide, x + Inches(0.1), y + Inches(0.58) + big_h, w - Inches(0.2),
             h - Inches(0.65) - big_h, sub, 12.5, GREY, align=PP_ALIGN.CENTER, spacing=0)


def new_slide(title, subtitle=None, notes=""):
    N["i"] += 1
    s = prs.slides.add_slide(BLANK)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.05))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.05), W, Inches(0.06))
    acc.fill.solid()
    acc.fill.fore_color.rgb = GOLD
    acc.line.fill.background()
    text(s, Inches(0.5), Inches(0.14), Inches(11.5), Inches(0.8), title, 30, WHITE, True)
    if subtitle:
        text(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.6), subtitle, 17, NAVY, True)
    text(s, Inches(12.2), Inches(7.0), Inches(1.0), Inches(0.4), str(N["i"]), 12, GREY,
         align=PP_ALIGN.RIGHT)
    text(s, Inches(0.5), Inches(7.0), Inches(8), Inches(0.4),
         "DERMA-GUARD  ·  VIT Chennai  ·  research prototype, not a medical device", 11, GREY)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def picture(slide, path, x, y, w=None, h=None):
    if not Path(path).exists():
        text(slide, x, y, Inches(4), Inches(0.5), f"[missing: {Path(path).name}]", 12, RED)
        return None
    from PIL import Image
    with Image.open(path) as im:
        r = im.height / im.width
    if w is not None and h is not None:          # fit inside the box, keep aspect ratio
        if w * r > h:
            w = int(h / r)
        else:
            h = int(w * r)
        return slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def table(slide, x, y, w, rows, col_w=None, size=14, header=NAVY, h_row=0.42):
    nr, nc = len(rows), len(rows[0])
    t = slide.shapes.add_table(nr, nc, x, y, w, Inches(h_row * nr)).table
    for j in range(nc):
        if col_w:
            t.columns[j].width = Inches(col_w[j])
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(size)
            p.font.name = FONT
            p.font.bold = i == 0
            p.font.color.rgb = WHITE if i == 0 else GREY
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            c.fill.solid()
            c.fill.fore_color.rgb = header if i == 0 else (LIGHT if i % 2 else WHITE)
    return t


# ================================================================ 1 title
s = prs.slides.add_slide(BLANK)
N["i"] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(1.6), Inches(0.09))
bar.fill.solid()
bar.fill.fore_color.rgb = GOLD
bar.line.fill.background()
text(s, Inches(0.8), Inches(1.4), Inches(11.7), Inches(1.2), "DERMA-GUARD", 60, WHITE, True)
text(s, Inches(0.8), Inches(2.55), Inches(11.7), Inches(1.4),
     "Evidence-Aware Specialist Consensus with Risk-Controlled Evidence Acquisition "
     "and Deferral for Multimodal Dermatology", 26, RGBColor(0xD9, 0xE2, 0xF3))
for k, (lab, c) in enumerate([("HIGH", C_HIGH), ("LOW", C_LOW), ("ACQUIRE", C_ASK), ("DEFER", C_DEFER)]):
    rect(s, Inches(0.8 + k * 1.75), Inches(3.75), Inches(1.6), Inches(0.45), c)
    text(s, Inches(0.8 + k * 1.75), Inches(3.78), Inches(1.6), Inches(0.4), lab, 14, WHITE, True,
         PP_ALIGN.CENTER)
rect(s, Inches(0.8), Inches(4.5), Inches(11.7), Inches(1.6), RGBColor(0x17, 0x3D, 0x60))
text(s, Inches(1.05), Inches(4.58), Inches(11.3), Inches(0.4), "TEAM", 13, GOLD, True)
text(s, Inches(1.05), Inches(4.92), Inches(11.3), Inches(0.5),
     "Rachit Agarwal (23BCE1577)   ·   Patel Kavya Vipulkumar   ·   Aarav Prawin", 22, WHITE, True)
text(s, Inches(1.05), Inches(5.5), Inches(11.3), Inches(0.5),
     "Guide: Dr. Sindhu Ravindran  ·  School of Computer Science and Engineering, VIT Chennai",
     16, RGBColor(0xD9, 0xE2, 0xF3))
text(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6),
     "An AI for skin lesions that knows when to answer, when to ask for more evidence, and "
     "when to hand over to a doctor.", 16, RGBColor(0xFF, 0xE6, 0x99))
s.notes_slide.notes_text_frame.text = (
    "Introduce the team and the one-line idea: we did not try to build an AI that is "
    "always right; we built one that knows when it should not answer on its own.")

# ================================================================ 2 problem
s = new_slide("The problem: skin-cancer AI always answers",
              "Accuracy on clean benchmarks is solved. Behaviour on real input is not.",
              "Explain the softmax problem in plain words: the probabilities must add up to one, so "
              "the model always picks something. Our stress test later shows blur makes the model MORE "
              "confident. In dermatology the costly error is a cancer confidently called harmless.")
text(s, Inches(0.6), Inches(1.9), Inches(6.4), Inches(4.8), [
    "A blurred photo, missing history, or a photo and history that disagree still get a "
    "confident label",
    "The confidence shown is rarely calibrated: '92%' does not mean 92% of such answers are right",
    "In our own stress test, blur LOWERED the model's uncertainty and RAISED its confident "
    "answers (46% to 50%)",
    "The costly error: a cancer called benign with confidence",
    "A careful dermatologist answers, looks closer, gives a provisional opinion, or refers"],
    19, bullet=True)
picture(s, FIG / "fig_evidential_example.png", Inches(7.2), Inches(2.1), w=Inches(5.8), h=Inches(4.4))

# ================================================================ 3 idea
s = new_slide("Our idea: four outcomes instead of one label",
              "Behave like a careful clinician, and measure the cost of doing so",
              "Walk through the four boxes. Stress that every outcome still carries the system's "
              "tentative opinion, differential and urgency, so even deferred patients get information.")
boxes = [("HIGH answer", "every safety check passed; calibrated diagnosis", C_HIGH),
         ("LOW answer", "best diagnosis, clearly marked; a doctor confirms", C_LOW),
         ("ACQUIRE", "ask for dermoscopy, missing details or a better photo", C_ASK),
         ("DEFER", "the evidence contradicts itself; a doctor decides", C_DEFER)]
for k, (t, d, c) in enumerate(boxes):
    x = Inches(0.6 + k * 3.15)
    rect(s, x, Inches(2.1), Inches(2.95), Inches(2.4), c)
    text(s, x + Inches(0.15), Inches(2.3), Inches(2.65), Inches(0.6), t, 24, WHITE, True, PP_ALIGN.CENTER)
    text(s, x + Inches(0.15), Inches(3.1), Inches(2.65), Inches(1.3), d, 16, WHITE, align=PP_ALIGN.CENTER)
text(s, Inches(0.6), Inches(4.9), Inches(12.2), Inches(2.0), [
    "Six specialists give opinions; a 'Devil's Advocate' decides how much to trust each, case by case",
    "A calibrated, cost-aware policy picks the outcome, with a certified bound on missed cancers",
    "A care platform turns outcomes into a workflow: patient app, doctor queue, messages, follow-ups"],
    18, bullet=True)

# ================================================================ 3b base papers
s = new_slide("Base papers: what we build on",
              "HOT AI for the task, Learn-then-Test for the guarantee",
              "Name the two base papers clearly. HOT AI showed an AI can ask for dermoscopy; we add "
              "costs, a defer action, a LOW tier and a certified bound, and we built a HOT-style "
              "baseline. Learn-then-Test is used exactly as published, on the missed-cancer rate. "
              "Our own earlier drafts (old1, old2) are where the team's ideas came from, not base papers.")
table(s, Inches(0.4), Inches(1.9), Inches(12.5), [
    ["", "Base paper 1: HOT AI (Mehta et al., JEADV 2025)", "Base paper 2: Learn-then-Test (Angelopoulos et al., 2025)"],
    ["What it does", "Recommends a dermoscopic image when the photo is not enough; OOD alert",
     "Chooses thresholds so an error rate stays below a target, with a written guarantee"],
    ["What we add", "Costs for asking and deferring; DEFER and LOW tiers; asks for details and better photos; "
     "certified bound; public data; stress and external tests",
     "Applied to the error that matters most: cancers called benign among confident answers"],
    ["How we compared", "HOT-style rule: ask the most uncertain, same rate, same probabilities",
     "Certified at both stages; checked on PAD-UFES-20 and Derm7pt"],
    ["Result", "It picks whom to ask better (AE +0.26 vs +0.07); but answering everyone after it misses "
     "5.6% vs our 3.1%", "Target 5% held: MILK10k 3.1%, PAD-UFES-20 0.0%, Derm7pt 3.8%"]],
    [2.3, 5.1, 5.1], 14, h_row=0.8)
text(s, Inches(0.4), Inches(6.3), Inches(12.5), Inches(0.6),
     "Also compared: learning to defer (4.9% missed), confidence threshold (3.8%), AUGRC, "
     "CAMA-style random/oracle acquisition, MILK10k reader study.", 15, NAVY)

# ================================================================ 4 data
s = new_slide("Data and an honest protocol",
              "MILK10k: the only large dataset with a clinical photo AND dermoscopy for every lesion",
              "Explain why MILK10k: to measure the value of asking for dermoscopy we need the real "
              "dermoscopic image of the same lesion. Split by lesion, validation decides everything, "
              "test scored once. PAD-UFES-20 and Derm7pt are external tests, never trained on.")
picture(s, FIG / "fig_class_distribution.png", Inches(0.4), Inches(1.9), w=Inches(7.4), h=Inches(3.6))
text(s, Inches(8.0), Inches(1.9), Inches(5.0), Inches(5.0), [
    "5,240 lesions, 11 diagnoses, 72% malignant, 96% biopsy-confirmed",
    "Split BY LESION: 3,667 train / 786 validation / 787 test",
    "Validation chooses everything; test scored once",
    "Cross-fitting: fusion is fitted on 4/5 of validation, scored on 1/5",
    "External tests, no retraining: PAD-UFES-20 (2,298 phone photos) and Derm7pt (1,003 "
    "lesions with dermoscopy)"], 17, bullet=True)
text(s, Inches(0.5), Inches(5.7), Inches(7.4), Inches(1.0),
     "Imbalanced: BCC is 48% of lesions; four classes have under 10 test lesions, so we report "
     "macro-F1 and balanced accuracy, not just accuracy.", 15, NAVY)

# ================================================================ 5 architecture
s = new_slide("System architecture", None,
              "Left to right: evidence, six specialists, the Devil's Advocate profile, the fusion "
              "rule chosen on validation, the gatekeeper, the four outcomes. The dashed line is the "
              "request for more evidence: the case is reassessed when it arrives.")
picture(s, FIG / "fig_architecture.png", Inches(0.5), Inches(1.3), w=Inches(12.3), h=Inches(5.6))

# ================================================================ 6 specialists + DA
s = new_slide("Six specialists and the Devil's Advocate",
              "Who sees what, and how much each is trusted for each diagnosis",
              "S1-S3 come from one evidential network trained on Kaggle. S4 is our old project's "
              "Model A. S5 and S6 reuse the learned embeddings. The heat map is the measured trust "
              "table; the old project typed these numbers in by hand (melanoma 0.96 assumed, 0.52 measured).")
table(s, Inches(0.4), Inches(1.9), Inches(6.6), [
    ["Specialist", "Acc.", "Macro-F1", "AUROC"],
    ["S1 photo head (evidential)", "0.595", "0.264", "0.787"],
    ["S2 dermoscopy head", "0.705", "0.400", "0.854"],
    ["S3 details head", "0.600", "0.246", "0.749"],
    ["S4 old Model A (PAD-UFES)", "0.219", "0.077", "0.553"],
    ["S5 similar-case vote", "0.719", "0.431", "0.866"],
    ["S6 embedding classifier", "0.700", "0.440", "0.875"]], [3.3, 1.1, 1.1, 1.1], 14)
text(s, Inches(0.4), Inches(5.2), Inches(6.6), Inches(1.6), [
    "Devil's Advocate signals: quality, learned reliability, uncertainty, confidence, margin, "
    "agreement, track record for the suspected diagnosis, consistency, missingness, OOD distance"],
    15, NAVY)
picture(s, FIG / "fig_trust_table.png", Inches(7.3), Inches(1.75), w=Inches(5.8), h=Inches(5.1))

# ================================================================ 7 fusion
s = new_slide("Fusion: seven rules, compared fairly",
              "The learned gate (Devil's Advocate v2) wins on validation; photo-only macro-F1 0.286 -> 0.377",
              "Grey bars are out-of-fold validation scores, used to choose; blue are test. Point out "
              "that stacking is the worst learned rule, and that the oracle shows choosing the right "
              "specialist per case is the real bottleneck (0.59 macro-F1).")
picture(s, FIG / "fig_fusion_rules.png", Inches(0.4), Inches(1.85), w=Inches(8.8), h=Inches(5.0))
card(s, Inches(9.45), Inches(1.9), Inches(3.6), Inches(1.55), "0.286 → 0.377",
     "photo-only macro-F1", "network alone → evidence-aware team", C_HIGH, 26)
card(s, Inches(9.45), Inches(3.55), Inches(3.6), Inches(1.55), "0.885",
     "cancer AUROC with dermoscopy", "accuracy 0.731, best of all rules", C_LOW, 26)
card(s, Inches(9.45), Inches(5.2), Inches(3.6), Inches(1.55), "0.594",
     "oracle ceiling (macro-F1)", "choosing whom to trust is the bottleneck", C_ASK, 26)

# ================================================================ 8 decision
s = new_slide("The decision policy",
              "Calibrated gates + a cost matrix; a missed cancer costs 5x a false alarm",
              "Read the flowchart top to bottom. The Learn-then-Test gate is what gives the "
              "certified bound: at most 5% of answered cancers missed, with 90% probability, if "
              "future patients resemble the calibration patients.")
picture(s, FIG / "fig_decision_flow.png", Inches(0.4), Inches(1.8), w=Inches(7.8), h=Inches(5.1))
text(s, Inches(8.4), Inches(1.9), Inches(4.7), Inches(5.0), [
    "Dempster-Shafer conflict separates 'contradiction' (defer) from 'not enough evidence' (ask)",
    "Conformal set over benign/malignant blocks false reassurance",
    "Learn-then-Test certifies the missed-cancer rate (alpha 5%, delta 10%)",
    "Quality, out-of-distribution and missing-details gates can only make it MORE careful",
    "Dermoscopy is requested only when its predicted value beats its cost"], 15, bullet=True)

# ================================================================ 9 results
s = new_slide("Results: what 787 unseen patients receive",
              "9 in 10 get an answer; 57% a confident one that misses 3.1% of cancers",
              "Emphasise the comparison at the SAME coverage on the SAME probabilities. The learned "
              "rejector is significantly less safe (p = 0.011); the confidence threshold is close "
              "(not significant) but, as the next slide shows, its safety is not stable.")
outs = [("57.1%", "HIGH answer", "77.1% accurate · 3.1% of cancers missed", C_HIGH),
        ("32.0%", "LOW answer", "a doctor confirms · true diagnosis in top-3 for 86%", C_LOW),
        ("29.5%", "ACQUIRE dermoscopy", "asked before deciding, then decided again", C_ASK),
        ("10.9%", "DEFER", "evidence contradicts itself · a doctor decides", C_DEFER)]
for k, (big, lab, sub_, c) in enumerate(outs):
    card(s, Inches(0.4 + k * 3.15), Inches(1.85), Inches(2.95), Inches(2.0), big, lab, sub_, c, 36)
table(s, Inches(7.0), Inches(4.1), Inches(6.0), [
    ["Same 57.1% coverage", "Accuracy", "Missed cancers"],
    ["DERMA-GUARD HIGH tier", "77.1%", "3.1%"],
    ["Confidence threshold", "81.5%", "3.8% (n.s.)"],
    ["Learned rejector (L2D)", "76.4%", "4.9% (p = 0.011)"],
    ["Always answer (photo)", "65.7%", "7.2%"]], [2.8, 1.4, 1.8], 13, h_row=0.4)
text(s, Inches(0.4), Inches(4.1), Inches(6.4), Inches(2.8), [
    [("89.1% answered (HIGH + LOW), 67.2% accurate overall", True, NAVY)],
    "Learn-then-Test certified at both stages (λ = 0.218 / 0.276)",
    "Displayed confidence is calibrated: 76% shown, 73% actual",
    "Against us: the threshold is more accurate; always answering is cheaper; HIGH answers "
    "raise more false alarms (50%)"], 15, bullet=True)

# ================================================================ 10 key finding
s = new_slide("Key finding: a threshold is only as safe as its probabilities",
              "Same policy on four fusion back-ends: DERMA-GUARD 3.0-3.5%, threshold 3.5-9.3%",
              "This is the most important result. Nobody deploying a threshold knows in advance "
              "whether their probabilities are the good kind or the bad kind. The calibrated policy "
              "gives the same certified miss rate either way.")
picture(s, FIG / "fig_backend_robustness.png", Inches(0.4), Inches(1.85), w=Inches(6.4), h=Inches(4.8))
picture(s, FIG / "fig_risk_coverage.png", Inches(7.0), Inches(1.85), w=Inches(6.0), h=Inches(4.8))

# ================================================================ 11 stress
s = new_slide("Robustness: DERMA-STRESS",
              "Uncertainty alone misses damaged photos; explicit checks catch all of them",
              "Blur makes the model more confident. Dark and occluded photos keep confident answers "
              "that miss 12-14% of cancers. Our quality checks (blur, light, grain, covered centre) and "
              "the request for missing details remove every confident answer on damaged input. "
              "Price: 8% of clean photos are flagged too.")
picture(s, FIG / "fig_stress.png", Inches(0.4), Inches(1.85), w=Inches(12.5), h=Inches(4.3))
text(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.7),
     "Blur: ignorance 0.679 -> 0.605 (fooled)   ·   with the checks: 0 confident answers on any damaged input",
     16, NAVY, True)

# ================================================================ 12 external
s = new_slide("External tests, without retraining",
              "PAD-UFES-20 (Brazil, phone photos) and Derm7pt (Italy, with dermoscopy)",
              "The system notices the shift: confident answers drop from 57% to 11-13%. Safety holds "
              "(0% and 3.8% missed). On Derm7pt asking for dermoscopy lifts accuracy 0.40 -> 0.48. "
              "But the 11-class diagnosis does not transfer. We did not train on these datasets on "
              "purpose: they measure what a deployment would meet.")
picture(s, FIG / "fig_shift.png", Inches(0.4), Inches(1.85), w=Inches(6.6), h=Inches(4.6))
table(s, Inches(7.2), Inches(1.95), Inches(5.8), [
    ["", "PAD-UFES-20", "Derm7pt"],
    ["Confident answers", "11.4%", "13.0%"],
    ["Missed cancers (HIGH)", "0.0%", "3.8%"],
    ["Cancer AUROC", "0.840", "0.738"],
    ["Dermoscopy requested", "28% (n/a)", "39%"],
    ["Accuracy, all cases", "0.367", "0.40 -> 0.48"]], [2.5, 1.65, 1.65], 14)
text(s, Inches(7.2), Inches(5.0), Inches(5.8), Inches(1.6),
     "Safe, not accurate, on a new population: the safety decision transfers, the fine "
     "diagnosis does not.", 16, NAVY, True)

# ================================================================ 13 old project
s = new_slide("Where we started: the old project, re-evaluated",
              "Its 84-90% were not real; its ideas were good",
              "Be factual and calm. The old split was by photo (43% of test photos showed lesions "
              "already in training) and the evaluation code adjusted predictions toward hard-coded "
              "target accuracies. Re-scored honestly: 57-75%. We kept its ideas (specialist team, "
              "Devil's Advocate, diagnosis-adaptive trust) and fused its seven models with our gate: "
              "72.7% vs 59.8% for averaging.")
picture(s, FIG / "fig_legacy_models.png", Inches(0.4), Inches(1.85), w=Inches(7.6), h=Inches(4.2))
text(s, Inches(8.2), Inches(1.95), Inches(4.9), Inches(5.0), [
    "Split by photo: 43% of test photos from lesions seen in training",
    "Evaluation flipped predictions to hit preset accuracy targets",
    "Honest scores on unseen patients: 57-75%",
    "Our learned gate over the old A-G: 72.7% (average: 59.8%)",
    "Kept: specialist team, Devil's Advocate, diagnosis-adaptive trust, Model A as S4"],
    15, bullet=True)
text(s, Inches(0.5), Inches(6.2), Inches(7.6), Inches(0.6),
     "Our own earlier drafts (old1, old2) had no measured results: placeholders and a plan.", 14, NAVY)

# ================================================================ 14 platform
s = new_slide("The platform: try every system yourself",
              "Patient app, doctor queue, explanations, compare-systems view, stress lab",
              "If time allows, demo live: Patient view with the demo lesion and dermoscopy unticked "
              "to show the request; Compare systems to run one lesion through every system; Stress "
              "Lab with blur. Every assessment comes with a 'why this result' explanation.")
def cropped(name, frac):
    """Top part of a tall screenshot, so it is readable on a slide."""
    from PIL import Image
    src = SCR / name
    if not src.exists():
        return src
    dst = ROOT / "reports" / "screens" / f"slide_{name}"
    with Image.open(src) as im:
        im.crop((260, 0, im.width, int(im.height * frac))).save(dst)
    return dst


picture(s, cropped("patient.png", 0.62), Inches(0.4), Inches(1.85), w=Inches(6.2), h=Inches(5.0))
picture(s, cropped("compare.png", 0.62), Inches(6.8), Inches(1.85), w=Inches(6.2), h=Inches(5.0))

# ================================================================ 14b findings
s = new_slide("Our findings at a glance",
              "Evidence-aware fusion decides what we believe; four outcomes decide what we do",
              "Summarise in one breath: trusting specialists case by case improves diagnosis; the four "
              "outcomes keep the confident answers safe; the safety is stable across models and "
              "datasets; damaged photos never get a confident answer; and asking for evidence pays.")
finds = [("+32%", "Evidence-aware fusion", "photo-only macro-F1 0.286 → 0.377; learned gate beats averaging and stacking", C_HIGH),
         ("3.1%", "Missed cancers, HIGH tier", "certified ≤ 5% by Learn-then-Test; learned rejector 4.9% (p = 0.011)", C_LOW),
         ("3.0–3.5%", "Stable on 4 fusion back-ends", "a confidence threshold swings 3.5–9.3% on the same models", C_ASK),
         ("0", "Confident answers on damaged photos", "blur fooled the uncertainty; the quality checks catch every case", C_DEFER),
         ("0.0% / 3.8%", "Safety on new datasets", "PAD-UFES-20 / Derm7pt missed cancers, no retraining", NAVY),
         ("0.40 → 0.48", "Asking for dermoscopy pays", "Derm7pt accuracy once the requested image arrives", GREY)]
for k, (big, lab, sub_, c) in enumerate(finds):
    r_, c_ = divmod(k, 3)
    card(s, Inches(0.4 + c_ * 4.2), Inches(1.85 + r_ * 2.5), Inches(4.0), Inches(2.3), big, lab, sub_, c, 34)

# ================================================================ 15 limits
s = new_slide("Limitations and what comes next", None,
              "Say the limits yourself before the panel finds them. The fairness gap is real and "
              "needs more lighter-tone data. Future work: dermatology foundation backbone, more "
              "seeds, better specialist selection (the oracle shows the headroom).")
text(s, Inches(0.6), Inches(1.4), Inches(6.1), Inches(5.4), [
    [("Limitations", True, NAVY)],
    "One training seed; ImageNet ResNet50 backbone",
    "Lighter skin tones worse served (tone 2: 11.5% missed)",
    "Fine diagnosis does not transfer to new populations",
    "Costs are declared, not clinically derived",
    "Guarantees assume future patients resemble calibration patients",
    "Not clinically validated - a research prototype"], 17, bullet=False, spacing=10)
text(s, Inches(6.9), Inches(1.4), Inches(6.1), Inches(5.4), [
    [("Next steps", True, NAVY)],
    "Dermatology foundation model as backbone",
    "More seeds: pipeline ready (run_seeds.py)",
    "Better per-case specialist selection (oracle 0.59 vs 0.44)",
    "Group-wise calibration by skin tone",
    "Learned image-quality model",
    "Prospective study with dermatologists"], 17, spacing=10)

# ================================================================ 16 conclusion
s = new_slide("Conclusion", None,
              "End on the three take-aways and invite questions.")
text(s, Inches(0.7), Inches(1.5), Inches(12), Inches(4.6), [
    "DERMA-GUARD answers 9 in 10 patients, more than half confidently, missing about 1 cancer "
    "in 30 - and asks for more evidence or a doctor when it should",
    "Learning whom to trust from all the evidence improves photo-only diagnosis by a third",
    "Its safety is certified and stable across models; a confidence threshold's is not",
    "Every number is reproducible: 230 tests, a script per table, honest re-evaluation of our "
    "own earlier work"], 21, bullet=True, spacing=14)
rect(s, Inches(3.4), Inches(5.8), Inches(6.5), Inches(0.9), NAVY)
text(s, Inches(3.4), Inches(5.93), Inches(6.5), Inches(0.7), "Thank you  -  questions?", 26, WHITE,
     True, PP_ALIGN.CENTER)

try:
    prs.save(OUT)
except PermissionError:                     # open in PowerPoint
    OUT = OUT.with_name(OUT.stem + "_updated.pptx")
    prs.save(OUT)
    print("  the presentation is open elsewhere - wrote a copy instead")
print(f"saved -> {OUT}  ({len(prs.slides)} slides)")
